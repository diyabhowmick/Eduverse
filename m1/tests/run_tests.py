"""Test suite. Run with: python -m tests.run_tests

No pytest dependency, so it runs on a bare Ollama box.
"""

from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import traceback

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from eduverse_book2docx import mathtools, omml, structure, textclean  # noqa: E402
from eduverse_book2docx.config import Config                        # noqa: E402
from eduverse_book2docx.ledger import Ledger                        # noqa: E402
from eduverse_book2docx.ocr import parse_markdown, _strip_inline    # noqa: E402
from eduverse_book2docx.ollama_client import OllamaClient, _safe_json  # noqa: E402
from eduverse_book2docx.pipeline import run                         # noqa: E402
from eduverse_book2docx.schema import Block, Book, Page, Run        # noqa: E402
from eduverse_book2docx.writer import template                      # noqa: E402
from tests.make_hard_sample import build as build_hard              # noqa: E402
from tests.make_sample import build as build_sample                 # noqa: E402

PASS: list = []
FAIL: list = []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append((name, detail))


def group(fn):
    def wrapper():
        try:
            fn()
        except Exception:
            FAIL.append((fn.__name__, traceback.format_exc(limit=3)))
    wrapper.__name__ = fn.__name__
    return wrapper


# ---------------------------------------------------------------------------
@group
def test_omml():
    cases = {
        "n = \\frac{c}{v}": ["m:f", "m:num", "m:den"],
        "x^{2}": ["m:sSup"],
        "a_{1}": ["m:sSub"],
        "x_{i}^{2}": ["m:sSubSup"],
        "\\sqrt{2}": ["m:rad", "degHide"],
        "\\sqrt[3]{8}": ["m:rad", "m:deg"],
        "\\sum_{i=1}^{n} i": ["m:nary", "\u2211"],
        "\\int_{0}^{1} x": ["m:nary", "\u222b"],
        "\\begin{pmatrix} a & b \\\\ c & d \\end{pmatrix}": ["m:m", "m:mr"],
        "\\left( \\frac{a}{b} \\right)": ["m:d", "begChr"],
        "\\alpha + \\beta": ["\u03b1", "\u03b2"],
        "\\hat{v}": ["m:acc"],
        "\\text{speed}": ["m:nor"],
        "45\\degree": ["\u00b0"],
    }
    for tex, needles in cases.items():
        xml = omml.latex_to_omml(tex)
        for n in needles:
            check(f"omml {tex[:20]} has {n}", n in xml)
        check(f"omml {tex[:20]} declares Cambria", "Cambria Math" in xml)

    check("omml display wraps oMathPara",
          omml.latex_to_omml("x", display=True).startswith("<m:oMathPara"))
    for bad in ["\\undefinedcmd{x}", "", "   "]:
        try:
            omml.latex_to_omml(bad)
            check(f"omml rejects {bad!r}", False)
        except omml.Unsupported:
            check(f"omml rejects {bad!r}", True)


@group
def test_mathtools():
    ok, why = mathtools.validate("\\frac{a}{b}")
    check("validate accepts good latex", ok, why)
    for bad, label in [("\\frac{a}{b", "brace"),
                       ("\\left( x", "left right"),
                       ("\\begin{pmatrix} a", "env"),
                       ("\\bogus{x}", "unsupported")]:
        ok, _ = mathtools.validate(bad)
        check(f"validate rejects {label}", not ok)

    spans = mathtools.find_math_spans("value is $x^2$ here")
    check("find inline math", len(spans) == 1 and spans[0][2] == "x^2")
    spans = mathtools.find_math_spans("$$E = mc^2$$")
    check("find display math", len(spans) == 1 and spans[0][3] is True)

    check("formula line detected", mathtools.looks_like_formula("v = u + at"))
    check("prose not a formula",
          not mathtools.looks_like_formula(
              "The speed of light in a vacuum is constant everywhere."))

    check("unicode superscript to latex",
          "^{2}" in mathtools.unicode_to_latex("x\u00b2"))
    check("times to latex", "\\times" in mathtools.unicode_to_latex("2\u00d73"))

    check("chemistry not maths", mathtools.is_chemistry("H2SO4"))
    check("equation not chemistry", not mathtools.is_chemistry("a = b + c"))
    runs = mathtools.split_chem_runs("H2O")
    check("chem subscripts", any(k == "sub" for _, k in runs))

    latex, why = mathtools.prepare("\\frac{a}{b")
    check("prepare fails closed", latex is None and why)


@group
def test_structure_sizes():
    cfg = Config()
    prof = structure.SizeProfile(cfg)
    for pg in range(50):
        prof.observe(11.0, 80, page=pg)          # body
    for pg in range(4):                          # a size must recur to count
        prof.observe(22.0, 40, page=pg)          # H1
        prof.observe(16.0, 60, page=pg)          # H2
        prof.observe(13.0, 60, page=pg)          # H3
    prof.finalise()
    check("body size found", abs(prof.body_size - 11.0) < 0.6, prof.body_size)
    check("three heading levels", len(prof.levels) == 3, prof.levels)
    check("H1 is level 1", prof.level_for(22.0) == 1)
    check("H2 is level 2", prof.level_for(16.0) == 2)
    check("body has no level", prof.level_for(11.0) is None)


@group
def test_structure_columns():
    cfg = Config()
    page = Page(page=1, width_pt=595, height_pt=842,
                margins_pt=[56, 56, 56, 56])
    blocks = []
    for i in range(4):
        blocks.append(Block(id=f"l{i}", page=1, type="paragraph",
                            bbox_pt=[56, 100 + i * 30, 250, 120 + i * 30],
                            text="left"))
        blocks.append(Block(id=f"r{i}", page=1, type="paragraph",
                            bbox_pt=[330, 100 + i * 30, 530, 120 + i * 30],
                            text="right"))
    n = structure.detect_columns(blocks, page, cfg)
    check("two columns detected", n == 2, str(n))

    single = [Block(id=f"s{i}", page=1, type="paragraph",
                    bbox_pt=[56, 100 + i * 30, 530, 120 + i * 30], text="x")
              for i in range(6)]
    check("single column detected",
          structure.detect_columns(single, page, cfg) == 1)

    structure.assign_columns(blocks, page, 2)
    page.columns = 2
    order = structure.reading_order(blocks, page, 2)
    check("left column read first",
          [b.id for b in order][:4] == ["l0", "l1", "l2", "l3"],
          str([b.id for b in order]))


@group
def test_structure_captions():
    cfg = Config()
    page = Page(page=1, width_pt=595, height_pt=842)
    fig = Block(id="f1", page=1, type="figure", bbox_pt=[100, 200, 400, 400])
    cap = Block(id="c1", page=1, type="paragraph",
                bbox_pt=[100, 410, 400, 424], text="Fig 2.4 A mirror")
    other = Block(id="p1", page=1, type="paragraph",
                  bbox_pt=[56, 600, 500, 640], text="Ordinary sentence here.")
    page.blocks = [fig, cap, other]
    structure.bind_captions(page, cfg)
    check("caption bound", cap.type == "caption" and cap.caption_of == "f1")
    check("prose left alone", other.type == "paragraph")
    ordered = structure.order_with_captions(page.blocks)
    check("caption follows figure",
          [b.id for b in ordered].index("c1")
          == [b.id for b in ordered].index("f1") + 1)


@group
def test_structure_runs():
    runs = [Run(s=0, e=5, bold=True), Run(s=5, e=9, bold=True),
            Run(s=9, e=12, italic=True), Run(s=12, e=12)]
    merged = structure.merge_runs(runs)
    check("adjacent runs merge", len(merged) == 2, str(merged))
    check("merged span correct", merged[0].s == 0 and merged[0].e == 9)


@group
def test_textclean():
    dirty = "a\x00b\x07c\x1fd"
    cleaned = textclean.clean(dirty)
    check("control chars replaced", "\x00" not in cleaned and "\x07" not in cleaned)
    check("clean is length preserving", len(cleaned) == len(dirty),
          f"{len(dirty)} -> {len(cleaned)}")
    check("tabs and newlines kept",
          textclean.clean("a\tb\nc") == "a\tb\nc")
    check("breaks dropped on request",
          textclean.clean("a\tb\nc", keep_breaks=False) == "a b c")
    check("private use replaced", textclean.clean("x\ue000y") == "x y")
    check("normal text untouched",
          textclean.clean("Ohm\u2019s law \u03a9 \u00b0C") == "Ohm\u2019s law \u03a9 \u00b0C")
    check("illegal detected", textclean.had_illegal("a\x0bb"))
    check("clean text not flagged", not textclean.had_illegal("normal text"))
    check("broken ratio counts", textclean.broken_ratio("\x00\x00ab") == 0.5)
    check("broken ratio zero on clean", textclean.broken_ratio("abcd") == 0.0)

    # the writer must survive a block full of control characters
    import tempfile
    from eduverse_book2docx.writer import flow_docx
    from eduverse_book2docx.schema import Block, Book, Page, Run
    b = Book(book_id="x", source_file="f.pdf")
    pg = Page(page=1, width_pt=595, height_pt=842)
    pg.blocks.append(Block(id="b1", page=1, type="paragraph",
                           bbox_pt=[56, 56, 500, 70],
                           text="bad\x00chars\x07here",
                           runs=[Run(s=0, e=3, bold=True)]))
    pg.blocks.append(Block(id="b2", page=1, type="heading", level=1,
                           bbox_pt=[56, 80, 500, 100], text="head\x1fing"))
    b.pages.append(pg)
    with tempfile.TemporaryDirectory() as d:
        path, stats = flow_docx.write(b, d, Config(), Ledger())
        from docx import Document
        doc = Document(path)
        txt = " ".join(p.text for p in doc.paragraphs)
        check("docx written despite control chars", os.path.exists(path))
        check("control chars gone from docx", "\x00" not in txt and "\x07" not in txt)
        check("surrounding text kept", "chars" in txt and "head" in txt)


@group
def test_ledger():
    led = Ledger()
    led.detect(10)
    led.emit(7)
    for i in range(3):
        led.drop(f"b{i}", 1, "empty")
    check("ledger balances", led.balanced())
    led.detect()
    check("ledger catches loss", not led.balanced())
    check("imbalance reported", led.imbalance() == 1)
    check("reasons counted", led.by_reason()["empty"] == 3)


@group
def test_markdown():
    md = """# Chapter one
Some **bold** text and *italic* text
that wraps onto a second line.

## 1.1 A section
- first bullet
- second bullet

| a | b |
|---|---|
| 1 | 2 |

<<FIGURE>>

$$x = 1$$
"""
    lines = parse_markdown(md)
    kinds = [l.kind for l in lines]
    check("heading parsed", kinds[0] == "heading" and lines[0].level == 1)
    check("wrapped paragraph joined",
          any(l.kind == "paragraph" and "second line" in l.text for l in lines))
    check("bold run captured",
          any(r.bold for l in lines for r in l.runs))
    check("subheading level", any(l.kind == "heading" and l.level == 2
                                  for l in lines))
    check("bullets parsed", kinds.count("list") == 2)
    check("table parsed", "table" in kinds)
    check("figure token parsed", "figure" in kinds)
    check("equation parsed", "equation" in kinds)

    text, runs = _strip_inline("a **b** c")
    check("markers stripped", text == "a b c")
    check("run offsets right", runs[0].s == 2 and runs[0].e == 3)


@group
def test_chat_endpoint_and_fallback():
    """Images must go over /api/chat, which carries them reliably, with
    /api/generate as the fallback. A chat call that comes back empty must fall
    through rather than be reported as the answer."""
    import json as _json
    from eduverse_book2docx import ollama_client as oc

    cfg = Config()
    cfg.offline = False

    calls = []

    class Resp:
        def __init__(self, payload):
            self._p = payload

        def raise_for_status(self):
            pass

        def json(self):
            return self._p

    def fake_post(url, json=None, timeout=None):
        calls.append((url.rsplit("/", 1)[-1], json))
        if url.endswith("/chat"):
            return Resp({"message": {"content": '{"desc": "from chat"}'}})
        return Resp({"response": '{"desc": "from generate"}'})

    real = oc.requests
    class R:
        post = staticmethod(fake_post)
    oc.requests = R
    try:
        c = oc.OllamaClient(cfg)
        d = c.describe_figure(b"png")
        check("chat is tried first", calls[0][0] == "chat", str(calls[0][0]))
        check("the image rides in the message",
              "images" in calls[0][1]["messages"][0], str(calls[0][1].keys()))
        check("thinking is turned off", calls[0][1].get("think") is False)
        check("json format requested", calls[0][1].get("format") == "json")
        check("chat answer used", d["desc"] == "from chat", str(d))
        check("endpoint recorded", c.calls[-1]["endpoint"] == "chat")

        # an empty chat reply must fall through to generate
        calls.clear()

        def empty_chat(url, json=None, timeout=None):
            calls.append((url.rsplit("/", 1)[-1], json))
            if url.endswith("/chat"):
                return Resp({"message": {"content": "   "}})
            return Resp({"response": '{"desc": "from generate"}'})

        R.post = staticmethod(empty_chat)
        c2 = oc.OllamaClient(cfg)
        d2 = c2.describe_figure(b"png")
        check("empty chat falls through",
              [x[0] for x in calls] == ["chat", "generate"], str(calls))
        check("generate answer used", d2["desc"] == "from generate", str(d2))
        check("the image rides in generate too",
              "images" in calls[1][1], str(calls[1][1].keys()))

        # both failing raises, and says so
        def dead(url, json=None, timeout=None):
            raise ConnectionError("refused")

        R.post = staticmethod(dead)
        c3 = oc.OllamaClient(cfg)
        try:
            c3.describe_figure(b"png")
            check("both endpoints failing raises", False)
        except oc.OllamaError as exc:
            check("both endpoints failing raises", "refused" in str(exc), str(exc))
        check("the failure is recorded", c3.calls[-1]["ok"] is False)
    finally:
        oc.requests = real


@group
def test_vision_probe():
    """A model that answers about nothing looks exactly like a model that will
    not describe. The health check rules that out with one tiny image."""
    from eduverse_book2docx import ollama_client as oc
    cfg = Config()
    cfg.offline = False

    class Base(oc.OllamaClient):
        reply = '{"shape": "circle", "colour": "red"}'

        def _post(self, model, prompt, images, num_ctx, want_json):
            self.saw_image = bool(images)
            return self.reply

    c = Base(cfg)
    r = c._vision_probe()
    check("a seeing model passes", r["sees_image"], str(r))
    check("the probe sends an image", c.saw_image)

    Base.reply = '{"shape": "I cannot see an image", "colour": "unknown"}'
    r = Base(cfg)._vision_probe()
    check("a blind model fails", not r["sees_image"], str(r))
    check("the reason is stated", "did not see" in r["note"], r["note"])

    class Dead(Base):
        def _post(self, *a, **k):
            raise oc.OllamaError("refused")

    r = Dead(cfg)._vision_probe()
    check("a failed probe fails", not r["sees_image"] and r["error"], str(r))


@group
def test_thinking_model_output():
    """qwen3.5 reasons out loud before it answers. That reasoning is not JSON
    and must never reach the parser."""
    d = _safe_json('<think>The figure shows a circle with two lines. Let me '
                   'describe it carefully.</think>'
                   '{"desc": "A circle", "fig_kind": "diagram"}',
                   {"desc": "", "fig_kind": "photo"})
    check("thinking block stripped", d["desc"] == "A circle", str(d))
    check("the answer still parses", d["fig_kind"] == "diagram")

    d = _safe_json("<think>reasoning that was cut off mid",
                   {"desc": "fallback"})
    check("unclosed thinking falls back", d["desc"] == "fallback", str(d))

    d = _safe_json('<THINK>upper case</THINK>```json\n{"desc": "ok"}\n```',
                   {"desc": ""})
    check("case and fences handled together", d["desc"] == "ok", str(d))

    d = _safe_json('{"desc": "no thinking here"}', {"desc": ""})
    check("plain json is untouched", d["desc"] == "no thinking here")


@group
def test_safe_json():
    d = _safe_json('```json\n{"a": 1}\n```', {"a": 0, "b": 2})
    check("fenced json parsed", d["a"] == 1 and d["b"] == 2)
    d = _safe_json("garbage", {"a": 0})
    check("garbage falls back", d["a"] == 0)
    d = _safe_json('prelude {"a": 3} tail', {"a": 0})
    check("embedded json found", d["a"] == 3)


@group
def test_offline_client():
    cfg = Config()
    cfg.offline = True
    c = OllamaClient(cfg)
    check("offline ocr returns nothing", c.ocr_page(b"") == "")
    check("offline split returns one part", len(c.split_composite(b"")) == 1)
    check("offline health ok", c.health()["ok"])


@group
def test_schema_roundtrip():
    b = Book(book_id="x", source_file="f.pdf")
    p = Page(page=1, width_pt=595, height_pt=842)
    blk = Block(id="b1", page=1, type="paragraph", bbox_pt=[1, 2, 3, 4],
                text="hello", runs=[Run(s=0, e=5, bold=True)])
    p.blocks.append(blk)
    b.pages.append(p)
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "book.json")
        b.save(path)
        b2 = Book.load(path)
    check("roundtrip block count", len(list(b2.blocks())) == 1)
    check("roundtrip run survives", b2.pages[0].blocks[0].runs[0].bold)
    check("roundtrip text survives", b2.pages[0].blocks[0].text == "hello")
    check("char total", b.char_total() == 5)


@group
def test_template():
    cfg = Config()
    with tempfile.TemporaryDirectory() as d:
        path = template.build(os.path.join(d, "t.docx"), cfg)
        from docx import Document
        doc = Document(path)
        names = {s.name for s in doc.styles}
        for want in ("EV H1", "EV H2", "EV Body", "EV Caption",
                     "EV Table Header", "EV Equation", "EV List Bullet"):
            check(f"style {want} exists", want in names)
        check("no stray paragraphs", len(doc.paragraphs) == 0)


# ---------------------------------------------------------------------------
@group
def test_pipeline_pdf():
    cfg = Config()
    cfg.offline = True
    with tempfile.TemporaryDirectory() as d:
        pdf = build_sample(os.path.join(d, "s.pdf"))
        out = os.path.join(d, "out")
        res = run(pdf, out, cfg)

        check("ledger balanced", res.ledger.balanced(),
              f"imbalance {res.ledger.imbalance()}")
        check("gate passed", res.passed, "; ".join(res.failures))
        check("four pages", len(res.book.pages) == 4)

        counts = res.book.counts()
        for t in ("heading", "paragraph", "list_item", "figure", "table",
                  "caption", "equation"):
            check(f"pdf produced {t}", counts.get(t, 0) > 0, str(counts))

        d_ = res.book.detected
        check("grade detected", d_.get("grade") == 8, str(d_.get("grade")))
        check("board detected", d_.get("board") is not None)
        check("subject detected", d_.get("subject") is not None)
        check("routes seen", set(d_["routes"]) >= {"native", "scanned"},
              str(d_["routes"]))

        check("figures written",
              len([f for f in os.listdir(os.path.join(out, "figures"))
                   if f.endswith(".png")]) >= 2)
        check("page rasters written",
              len(os.listdir(os.path.join(out, "pages"))) == 4)
        for f in ("book.json", "flow.docx", "report.html", "ledger.jsonl"):
            check(f"{f} produced", os.path.exists(os.path.join(out, f)))

        # bold survived into the docx
        from docx import Document
        doc = Document(os.path.join(out, "flow.docx"))
        styles = [p.style.name for p in doc.paragraphs]
        check("EV heading styles used",
              any(s.startswith("EV H") for s in styles), str(set(styles)))
        check("bullet style used", "EV List Bullet" in styles)
        check("tables written",
              len([t for t in doc.tables if len(t.columns) > 1]) == 2,
              str(len(doc.tables)))
        check("pictures written", len(doc.inline_shapes) == 2,
              str(len(doc.inline_shapes)))
        omath = doc.element.body.findall(
            ".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")
        check("equation is native OMML", len(omath) >= 1)
        alts = [el.get("descr") for el in doc.element.body.iter()
                if el.tag.endswith("}docPr") and el.get("descr")]
        check("alt text carries block id",
              any("eduverse:" in (a or "") for a in alts))

        # ledger file is readable and carries the invariant
        with open(os.path.join(out, "ledger.jsonl"), encoding="utf-8") as fh:
            summary = json.loads(fh.readline())["_summary"]
        check("ledger file balanced", summary["balanced"])
        check("every drop has a reason",
              all(json.loads(l).get("reason")
                  for l in open(os.path.join(out, "ledger.jsonl"),
                                encoding="utf-8").readlines()[1:]))


@group
def test_pipeline_docx_roundtrip():
    cfg = Config()
    cfg.offline = True
    with tempfile.TemporaryDirectory() as d:
        pdf = build_sample(os.path.join(d, "s.pdf"))
        out1 = os.path.join(d, "o1")
        r1 = run(pdf, out1, cfg)
        out2 = os.path.join(d, "o2")
        r2 = run(r1.files["flow_docx"], out2, cfg)

        check("docx path balanced", r2.ledger.balanced())
        check("docx path passed", r2.passed, "; ".join(r2.failures))
        c1, c2 = r1.book.counts(), r2.book.counts()
        check("headings survive the round trip",
              c2.get("heading", 0) == c1.get("heading", 0),
              f"{c1.get('heading')} -> {c2.get('heading')}")
        check("figures survive", c2.get("figure", 0) == c1.get("figure", 0))
        check("tables survive", c2.get("table", 0) == c1.get("table", 0))
        check("equation survives as OMML",
              r2.stats.get("equations", 0) >= 1, str(r2.stats))
        check("docx format detected", r2.book.detected["format"] == "docx")


@group
def test_layout_and_page_lock():
    cfg = Config()
    cfg.offline = True
    cfg.write_layout_docx = True
    cfg.page_lock = True
    with tempfile.TemporaryDirectory() as d:
        pdf = build_sample(os.path.join(d, "s.pdf"))
        out = os.path.join(d, "out")
        res = run(pdf, out, cfg)
        check("layout.docx produced", os.path.exists(res.files["layout_docx"]))
        from docx import Document
        doc = Document(res.files["layout_docx"])
        framed = doc.element.body.findall(
            ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}framePr")
        check("blocks are frame anchored", len(framed) > 3, str(len(framed)))


class _FakeClient(OllamaClient):
    """Stands in for Ollama so the description and text-only paths can be
    tested without the models."""
    ocr_markdown = ""
    describe_as = None

    def ocr_page(self, png_bytes):
        return self.ocr_markdown

    def describe_figure(self, png_bytes):
        return dict(self.describe_as or {})

    def split_composite(self, png_bytes):
        return [{"x0": 0.0, "y0": 0.0, "x1": 1.0, "y1": 1.0}]


def _run_with_fake(tmp, describe_as, ocr_markdown="", **opts):
    from eduverse_book2docx import pipeline
    cfg = Config()
    cfg.offline = False
    for k, v in opts.items():
        setattr(cfg, k, v)
    _FakeClient.describe_as = describe_as
    _FakeClient.ocr_markdown = ocr_markdown
    real = pipeline.OllamaClient
    pipeline.OllamaClient = _FakeClient
    try:
        pdf = build_sample(os.path.join(tmp, "s.pdf"))
        return pipeline.run(pdf, os.path.join(tmp, "out"), cfg)
    finally:
        pipeline.OllamaClient = real


@group
def test_figure_description_box():
    desc = {
        "fig_kind": "diagram", "text_only": False, "label": "Fig 2.4",
        "desc": "A plane mirror drawn as a horizontal hatched line. An "
                "incident ray arrives from the upper left at point O.",
        "entities": ["plane mirror", "incident ray", "normal"],
        "relationships": ["angle i equals angle r"],
        "data_points": [], "text_in_image": ["i", "r", "O", "N"],
    }
    with tempfile.TemporaryDirectory() as d:
        res = _run_with_fake(d, desc)
        counts = res.book.counts()
        check("description box created", counts.get("figure_note", 0) >= 1,
              str(counts))
        check("ledger still balanced", res.ledger.balanced(),
              f"imbalance {res.ledger.imbalance()}")
        check("gate passed with boxes", res.passed, "; ".join(res.failures))

        notes = [b for _, b in res.book.blocks() if b.type == "figure_note"]
        figs = [b for _, b in res.book.blocks() if b.type == "figure"]
        fig_notes = [n for n in notes if n.caption_of in {f.id for f in figs}]
        check("every figure has a box", len(fig_notes) == len(figs),
              f"{len(figs)} figures, {len(fig_notes)} boxes")
        check("box points at its figure",
              all(n.caption_of in {f.id for f in figs} for n in fig_notes))
        check("box carries the description",
              any("plane mirror" in n.text for n in fig_notes))
        check("box lists the parts",
              any("Parts:" in n.text for n in fig_notes))
        check("box lists the text in the figure",
              any("Text in figure:" in n.text for n in fig_notes))

        # ordering: figure, then its box, then the printed caption
        page = res.book.pages[0]
        ids = [b.type for b in page.blocks]
        fi = ids.index("figure")
        check("box sits directly under the figure",
              ids[fi + 1] == "figure_note", str(ids))
        check("printed caption still follows",
              "caption" in ids[fi + 1: fi + 4], str(ids))

        from docx import Document
        doc = Document(res.files["flow_docx"])
        cell_text = " ".join(c.text for t in doc.tables
                             for row in t.rows for c in row.cells)
        check("box rendered into the docx", "plane mirror" in cell_text)
        check("box is labelled", "Image description" in cell_text)
        check("box counted", res.stats.get("figure_notes", 0) >= 1)
        check("alt text still carries it too",
              any("plane mirror" in (el.get("descr") or "")
                  for el in doc.element.body.iter()
                  if el.tag.endswith("}docPr")))


@group
def test_text_only_region_becomes_text():
    md = ("## Activity 2\n"
          "Hold a **plane mirror** upright on a table.\n"
          "- Shine a torch at the mirror.\n"
          "- Mark where the beam lands.\n")
    desc = {"fig_kind": "text_banner", "text_only": True, "label": "",
            "desc": "A block of printed text.", "entities": [],
            "relationships": [], "data_points": [], "text_in_image": []}
    with tempfile.TemporaryDirectory() as d:
        res = _run_with_fake(d, desc, ocr_markdown=md)
        counts = res.book.counts()
        check("no picture kept for a text region",
              counts.get("figure", 0) == 0, str(counts))
        check("no description box for the converted text",
              not any(b.type == "figure_note" and b.block_kind not in
                      ("equation", "table")
                      for _, b in res.book.blocks()), str(counts))
        check("ledger balanced after conversion", res.ledger.balanced(),
              f"imbalance {res.ledger.imbalance()}")

        texts = [b.text for _, b in res.book.blocks()]
        check("region text was published",
              any("plane mirror" in t for t in texts), str(texts)[:300])
        check("heading inside the region kept its level",
              any(b.type == "heading" and "Activity" in b.text
                  for _, b in res.book.blocks()))
        check("bullets inside the region became list items",
              sum(1 for _, b in res.book.blocks()
                  if b.type == "list_item" and "torch" in b.text
                  or b.type == "list_item" and "beam" in b.text) >= 1)
        check("bold inside the region survived",
              any(r.bold for _, b in res.book.blocks() for r in b.runs
                  if "plane mirror" in b.text))
        check("blocks marked as coming from a text region",
              any("from_text_only_region" in b.flags
                  for _, b in res.book.blocks()))
        check("drop reason recorded",
              any("published as text" in (e.get("detail") or "")
                  for e in res.ledger.entries))

        from docx import Document
        doc = Document(res.files["flow_docx"])
        body = " ".join(p.text for p in doc.paragraphs)
        check("text landed in the docx", "plane mirror" in body)
        check("no picture in the docx", len(doc.inline_shapes) == 0,
              str(len(doc.inline_shapes)))


@group
def test_description_box_can_be_switched_off():
    desc = {"fig_kind": "photo", "text_only": False, "label": "",
            "desc": "A photograph of a mirror.", "entities": [],
            "relationships": [], "data_points": [], "text_in_image": []}
    with tempfile.TemporaryDirectory() as d:
        res = _run_with_fake(d, desc, figure_description_box=False)
        check("no boxes when switched off",
              not any(b.type == "figure_note" and b.block_kind not in
                      ("equation", "table") for _, b in res.book.blocks()))
        check("figures still kept", res.book.counts().get("figure", 0) >= 1)
        check("ledger balanced", res.ledger.balanced())


@group
def test_tables():
    from eduverse_book2docx import tables as tmod
    cfg = Config()
    cfg.offline = True
    with tempfile.TemporaryDirectory() as d:
        pdf = build_sample(os.path.join(d, "s.pdf"))
        out = os.path.join(d, "out")
        res = run(pdf, out, cfg)

        tbls = [b for _, b in res.book.blocks() if b.type == "table"]
        check("both tables found", len(tbls) == 2, str(len(tbls)))

        simple = next((t for t in tbls if t.table.cols == 3), None)
        check("simple table shape",
              simple is not None and simple.table.rows == 4,
              str(simple.table.rows) if simple else "missing")
        if simple:
            g = {(c.row, c.col): c.text for c in simple.table.cells}
            check("header cell right", g.get((0, 0)) == "Medium", str(g.get((0, 0))))
            check("data cell right", g.get((2, 2)) == "225000", str(g.get((2, 2))))
            check("header row detected", simple.table.header_rows == 1)
            check("bold header kept",
                  any(c.runs and all(r.bold for r in c.runs)
                      for c in simple.table.cells if c.row == 0))
            check("ruling lines detected",
                  simple.table.has_v_rules and simple.table.has_h_rules)
            check("column widths are real",
                  all(w > 20 for w in simple.table.col_widths_pt),
                  str(simple.table.col_widths_pt))
            check("widths sum near the table width",
                  abs(sum(simple.table.col_widths_pt) - simple.width) < 12,
                  f"{sum(simple.table.col_widths_pt):.0f} vs {simple.width:.0f}")

        merged = next((t for t in tbls if t.table.cols == 4), None)
        check("merged table shape",
              merged is not None and merged.table.rows == 6,
              str(merged.table.rows) if merged else "missing")
        if merged:
            g = {(c.row, c.col): c for c in merged.table.cells}
            check("merge detected from missing rules",
                  g[(0, 1)].col_span == 3, str(g[(0, 1)].col_span))
            check("cell under the merge is a sub header",
                  g[(1, 1)].text == "Index", g[(1, 1)].text)
            check("last data row survived",
                  g[(5, 0)].text == "Glass", g[(5, 0)].text)
            check("no cell text is tab joined",
                  all("\t" not in c.text for c in merged.table.cells))

        # no table content leaked into the text flow
        texts = " ".join(b.text for _, b in res.book.blocks()
                         if b.type in ("paragraph", "heading", "list_item"))
        # "Refractive index" is also a real heading on that page, so check
        # values that only ever appear inside a table
        check("table cells not duplicated as prose",
              "299700" not in texts and "Speed km per s" not in texts
              and "225000" not in texts,
              texts[-160:])

        # captions still bound
        caps = [b for _, b in res.book.blocks()
                if b.type == "caption" and b.text.lower().startswith("table")]
        check("table captions bound", len(caps) == 2, str(len(caps)))
        check("captions point at tables",
              all(c.caption_of in {t.id for t in tbls} for c in caps))

        # and the docx side
        from docx import Document
        doc = Document(res.files["flow_docx"])
        real = [t for t in doc.tables if len(t.columns) > 1]
        check("two tables in the docx", len(real) == 2, str(len(real)))
        t2 = real[1]
        check("docx grid is 6 by 4",
              len(t2.rows) == 6 and len(t2.columns) == 4,
              f"{len(t2.rows)}x{len(t2.columns)}")
        check("merged cell is merged in the docx",
              t2.cell(0, 1)._tc is t2.cell(0, 3)._tc)
        check("header row repeats across pages",
              t2.rows[0]._tr.find(
                  "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}trPr"
              ) is not None)
        widths = [c._tc.find(
            ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcW")
            for c in t2.rows[1].cells]
        check("every cell carries a width", all(w is not None for w in widths))

    check("summary reads cleanly",
          "ruled" in tmod.summary(tbls[0].table))


@group
def test_table_from_markdown():
    from eduverse_book2docx import tables as tmod
    cfg = Config()
    rows = [["Medium", "Index"], ["Water", "1.33"], ["Glass", "1.50"]]
    blk = tmod.from_markdown(rows, "t1", 1, [56, 100, 500, 180], cfg)
    check("markdown table built", blk is not None and blk.type == "table")
    check("markdown shape", blk.table.rows == 3 and blk.table.cols == 2)
    check("markdown header", blk.table.header_rows == 1)
    check("markdown widths split evenly",
          abs(sum(blk.table.col_widths_pt) - 444) < 2,
          str(blk.table.col_widths_pt))
    check("one row is not a table",
          tmod.from_markdown([["a", "b"]], "t2", 1, [0, 0, 10, 10], cfg) is None)
    check("one column is not a table",
          tmod.from_markdown([["a"], ["b"], ["c"]], "t3", 1,
                             [0, 0, 10, 10], cfg) is None)
    check("empty rows dropped",
          tmod.from_markdown([["Medium", "Index"], ["", ""], ["Water", "1.33"]],
                             "t4", 1, [0, 0, 100, 100], cfg).table.rows == 2)


@group
def test_real_book_shape():
    """The faults a real state textbook produced: every bold label became an
    H4, running heads and page numbers stayed in the text, and a boxed activity
    came out as one paragraph per line."""
    cfg = Config()
    cfg.offline = True
    with tempfile.TemporaryDirectory() as d:
        pdf = build_hard(os.path.join(d, "hard.pdf"))
        res = run(pdf, os.path.join(d, "out"), cfg)

        check("hard book balanced", res.ledger.balanced(),
              f"imbalance {res.ledger.imbalance()}")
        check("hard book gate passed", res.passed, "; ".join(res.failures))

        blocks = [b for _, b in res.book.blocks()]
        texts = [b.text.strip() for b in blocks if b.text]

        # running head, page numbers and side tab all gone
        check("running head dropped",
              not any(t == "Science 8" for t in texts), str(texts[:6]))
        check("page numbers dropped",
              not any(t.isdigit() and len(t) <= 2 for t in texts),
              str([t for t in texts if t.isdigit()]))
        check("side tab dropped", not any(t == "Unit 1" for t in texts))
        check("furniture count recorded",
              res.book.stats.get("furniture_dropped", 0) >= 6,
              str(res.book.stats.get("furniture_dropped")))

        # heading levels follow the book's own scale, nothing invented
        levels = sorted({b.level for b in blocks if b.type == "heading"})
        sizes = res.book.stats["size_profile"]["heading_sizes"]
        # a size based level never exceeds the sizes the book uses, but a
        # numbered heading may go deeper, because its own numbering says so
        numbered_depth = max(
            (b.text.split()[0].count(".") + 1 for b in blocks
             if b.type == "heading" and b.text[:1].isdigit()), default=1)
        check("no invented heading depth",
              max(levels) <= max(len(sizes), numbered_depth),
              f"{levels} vs {len(sizes)} sizes, numbered {numbered_depth}")
        check("h1 is the title", any(b.type == "heading" and b.level == 1
                                     and "MEASUREMENT" in b.text for b in blocks))
        check("numbered section is h2",
              any(b.type == "heading" and b.level == 2
                  and b.text.startswith("1.1 ") for b in blocks))
        check("sub section is h3",
              any(b.type == "heading" and b.level == 3
                  and b.text.startswith("1.1.1") for b in blocks))

        # bold body-sized labels stay paragraphs, but keep their bold
        labels = [b for b in blocks if b.text.strip() in
                  ("Learning Objectives", "Introduction", "Do you know?")]
        check("run-in labels found", len(labels) == 3, str(len(labels)))
        check("run-in labels are not headings",
              all(b.type != "heading" for b in labels),
              str([(b.text, b.type, b.level) for b in labels]))
        check("run-in labels keep their bold",
              all(b.runs and any(r.bold for r in b.runs) for b in labels))

        # the boxed activity is text in a panel, not a picture. The two real
        # figures on that page are the clock and the six part illustration.
        panel_figs = [b for b in blocks if b.type == "figure"
                      and b.x0 < 310 and b.y0 < 340]
        check("panel is not a figure", not panel_figs,
              str([b.id for b in panel_figs]))
        check("panel heading survived",
              any(b.text.strip() == "Activity 1" for b in blocks))
        check("panel blocks are marked",
              any(b.block_kind == "panel" for b in blocks))
        check("activity text is one paragraph, not one per line",
              any("using a ruler" in b.text and "Measure" in b.text
                  for b in blocks),
              str([b.text[:40] for b in blocks if "Measure" in b.text]))

        # nothing invented, nothing duplicated
        check("no empty paragraphs",
              all(b.text.strip() for b in blocks if b.type == "paragraph"))
        check("no duplicated paragraph",
              len(texts) == len(set(texts)),
              str([t for t in texts if texts.count(t) > 1][:3]))


@group
def test_fragment_joining():
    from eduverse_book2docx.structure import join_fragments
    cfg = Config()
    page = Page(page=1, width_pt=595, height_pt=842)

    def blk(i, text, y, x=56.0):
        return Block(id=f"b{i}", page=1, type="paragraph",
                     bbox_pt=[x, y, x + 200, y + 12], text=text)

    # a wrapped sentence: joins
    page.blocks = [blk(1, "Measure the length", 100),
                   blk(2, "and breadth of your", 113),
                   blk(3, "science book.", 126)]
    n = join_fragments(page, cfg)
    check("wrapped sentence joined", len(page.blocks) == 1 and n == 2,
          str([b.text for b in page.blocks]))
    check("joined text reads correctly",
          page.blocks[0].text == "Measure the length and breadth of your "
                                 "science book.", page.blocks[0].text)

    # a list of terms: must NOT join
    page2 = Page(page=1, width_pt=595, height_pt=842)
    page2.blocks = [blk(1, "Length", 100), blk(2, "Mass", 113),
                    blk(3, "Time", 126), blk(4, "Temperature", 139)]
    join_fragments(page2, cfg)
    check("list of terms left alone", len(page2.blocks) == 4,
          str([b.text for b in page2.blocks]))

    # a finished sentence followed by a new one: must NOT join
    page3 = Page(page=1, width_pt=595, height_pt=842)
    page3.blocks = [blk(1, "This is done.", 100),
                    blk(2, "then this follows", 113)]
    join_fragments(page3, cfg)
    check("finished sentence not joined", len(page3.blocks) == 2)

    # far apart vertically: must NOT join
    page4 = Page(page=1, width_pt=595, height_pt=842)
    page4.blocks = [blk(1, "an unfinished line", 100),
                    blk(2, "that is far below", 400)]
    join_fragments(page4, cfg)
    check("distant blocks not joined", len(page4.blocks) == 2)

    # runs survive the join with correct offsets
    page5 = Page(page=1, width_pt=595, height_pt=842)
    a = blk(1, "the speed of", 100)
    bb = blk(2, "light is constant", 113)
    bb.runs = [Run(s=0, e=5, bold=True)]
    page5.blocks = [a, bb]
    join_fragments(page5, cfg)
    joined = page5.blocks[0]
    check("run offset shifted on join",
          joined.runs and joined.text[joined.runs[0].s:joined.runs[0].e] == "light",
          str([(r.s, r.e) for r in joined.runs]) + " " + joined.text)


@group
def test_repeating_furniture():
    from eduverse_book2docx.structure import find_repeating_furniture
    from eduverse_book2docx.schema import Book as B

    book = B(book_id="x", source_file="f.pdf")
    for i in range(1, 9):
        pg = Page(page=i, width_pt=595, height_pt=842)
        pg.blocks = [
            Block(id=f"h{i}", page=i, type="paragraph",
                  bbox_pt=[56, 30, 200, 42], text="Science 8"),
            Block(id=f"n{i}", page=i, type="paragraph",
                  bbox_pt=[520, 800, 540, 812], text=str(i)),
            Block(id=f"b{i}", page=i, type="paragraph",
                  bbox_pt=[56, 200, 500, 300],
                  text=f"Real body text unique to page {i} about measurement."),
        ]
        book.pages.append(pg)
    reps = find_repeating_furniture(book, Config())
    check("running head detected as furniture", "science #" in reps, str(reps))
    check("body text not furniture",
          not any("real body" in k for k in reps), str(reps))

    led = Ledger()
    led.emit(24)
    from eduverse_book2docx.structure import drop_furniture
    removed = drop_furniture(book, led, Config())
    check("furniture removed", removed == 16, str(removed))
    left = [b.text for p in book.pages for b in p.blocks]
    check("only body text left", all("Real body" in t for t in left), str(left[:4]))
    check("emitted count corrected", led.emitted == 8, str(led.emitted))

    # a three page book is too short to judge repetition
    short = B(book_id="y", source_file="g.pdf")
    for i in range(1, 4):
        pg = Page(page=i, width_pt=595, height_pt=842)
        pg.blocks = [Block(id=f"h{i}", page=i, type="paragraph",
                           bbox_pt=[56, 30, 200, 42], text="Chapter One")]
        short.pages.append(pg)
    check("short book is not judged",
          find_repeating_furniture(short, Config()) == {})


@group
def test_svg_from_spec():
    from eduverse_book2docx import svgdraw
    chart = {"kind": "chart", "chart_type": "bar", "title": "Speed of light",
             "x_label": "Medium", "y_label": "km per s",
             "categories": ["Vacuum", "Water", "Glass"],
             "series": [{"name": "speed", "values": [300000, 225000, 200000]}]}
    svg = svgdraw.build(chart, 200, 150)
    check("chart drawn", svg is not None and svg.startswith("<svg"))
    check("bars drawn", svg.count("<rect") >= 4, str(svg.count("<rect")))
    check("axis label drawn", "km per s" in svg)
    check("categories drawn", "Vacuum" in svg and "Glass" in svg)
    check("chart is regenerable", svgdraw.is_regenerable(chart))

    diagram = {"kind": "diagram", "layout": "flow",
               "nodes": [{"id": "a", "label": "incident ray", "shape": "box"},
                         {"id": "b", "label": "mirror", "shape": "box"},
                         {"id": "c", "label": "reflected ray", "shape": "box"}],
               "edges": [{"from": "a", "to": "b", "label": "angle i"},
                         {"from": "b", "to": "c", "label": "angle r"}]}
    svg = svgdraw.build(diagram, 300, 200)
    check("diagram drawn", svg is not None and "incident ray" in svg)
    check("edges drawn", "marker-end" in svg)
    check("diagram is regenerable", svgdraw.is_regenerable(diagram))

    # not enough to redraw
    check("empty spec is not regenerable", not svgdraw.is_regenerable({}))
    check("one node is not regenerable",
          not svgdraw.is_regenerable({"nodes": [{"id": "a", "label": "x"}]}))
    check("chart with no readable values is not regenerable",
          not svgdraw.is_regenerable(
              {"chart_type": "bar",
               "series": [{"values": [None, None, None]}]}))
    check("unknown spec draws nothing",
          svgdraw.build({"kind": "photo"}, 100, 100) is None)
    check("labels are escaped",
          "&lt;b&gt;" in svgdraw.build(
              {"nodes": [{"id": "a", "label": "<b>x"},
                         {"id": "b", "label": "y"}]}, 200, 200))


@group
def test_coverage():
    from eduverse_book2docx import coverage as cov_mod
    from eduverse_book2docx.schema import Book as B, FigureMeta, TableCell, TableData

    book = B(book_id="x", source_file="f.pdf")
    pg = Page(page=1, width_pt=595, height_pt=842)
    pg.blocks = [
        Block(id="p1", page=1, type="paragraph", bbox_pt=[0, 0, 1, 1],
              text="readable prose"),
        Block(id="e1", page=1, type="equation", bbox_pt=[0, 0, 1, 1],
              text="x", latex="x^{2}", latex_source="vlm"),
        Block(id="e2", page=1, type="equation", bbox_pt=[0, 0, 1, 1],
              text="?", math_unreadable=True, image_ref="figures/e2.png"),
        Block(id="f1", page=1, type="figure", bbox_pt=[0, 0, 1, 1],
              figure=FigureMeta(desc="a diagram", regenerable=True,
                                svg_ref="figures/f1.svg")),
        Block(id="f2", page=1, type="figure", bbox_pt=[0, 0, 1, 1],
              figure=FigureMeta(desc="")),
        Block(id="t1", page=1, type="table", bbox_pt=[0, 0, 1, 1],
              source="vlm_table",
              table=TableData(rows=1, cols=1,
                              cells=[TableCell(row=0, col=0, text="v")])),
        Block(id="fur", page=1, type="page_furniture", bbox_pt=[0, 0, 1, 1],
              text="12"),
    ]
    book.pages.append(pg)
    c = cov_mod.measure(book)

    check("furniture not counted", c.total == 6, str(c.total))
    check("opaque counted", c.opaque == 2, str(c.opaque))
    check("coverage ratio", abs(c.ratio - 4 / 6) < 0.01, str(c.ratio))
    ids = {o["id"] for o in c.opaque_blocks}
    check("unreadable equation is opaque", "e2" in ids, str(ids))
    check("undescribed figure is opaque", "f2" in ids, str(ids))
    check("described figure is readable", "f1" not in ids)
    check("latex equation is readable", "e1" not in ids)
    check("equation source tracked", c.equations_from_vlm == 1)
    check("svg counted", c.vector == 1)
    check("regenerable counted", c.regenerable == 1)
    check("table from image counted", c.tables_from_image == 1)
    check("worst page listed", cov_mod.worst_pages(c)[0] == (1, 2))
    check("summary is serialisable",
          json.loads(json.dumps(c.summary()))["opaque"] == 2)


@group
def test_digital_book():
    desc = {"fig_kind": "chart", "text_only": False, "label": "Fig 2.4",
            "desc": "Bar chart of the speed of light.",
            "entities": ["vacuum", "glass"], "relationships": ["denser is slower"],
            "data_points": ["vacuum 300000"], "text_in_image": ["km/s"]}

    class SpecClient(_FakeClient):
        def figure_spec(self, png, kind):
            return {"kind": "chart", "chart_type": "bar",
                    "title": "Speed of light", "y_label": "km per s",
                    "categories": ["Vacuum", "Water", "Glass"],
                    "series": [{"name": "speed",
                                "values": [300000, 225000, 200000]}]}

        def read_table(self, png):
            return {"is_table": False}

        def read_math(self, png):
            return {"latex": "", "confident": False}

    from eduverse_book2docx import pipeline as pl
    cfg = Config()
    cfg.offline = False
    SpecClient.describe_as = desc
    SpecClient.ocr_markdown = ""
    real = pl.OllamaClient
    pl.OllamaClient = SpecClient
    try:
        with tempfile.TemporaryDirectory() as d:
            pdf = build_sample(os.path.join(d, "s.pdf"))
            res = pl.run(pdf, os.path.join(d, "out"), cfg)
            cs = res.coverage.summary()

            check("nothing opaque", cs["opaque"] == 0,
                  str(res.coverage.opaque_blocks))
            check("full coverage", cs["coverage"] == 1.0, str(cs["coverage"]))
            check("figures regenerable",
                  cs["figures_regenerable"] == cs["figures"], str(cs))
            check("svg written", cs["figures_with_svg"] >= 1)

            path = res.files.get("digital")
            check("digital book written", path and os.path.exists(path))
            h = open(path, encoding="utf-8").read()

            check("figures are inline svg", "<svg" in h)
            check("maths is mathml", "<math" in h)
            check("tables are real tables", "<table" in h and "<th" in h)
            check("lists are real lists", "<li" in h)
            check("bullet glyph stripped",
                  "<li id=" in h and "\u2022" not in h.split("<li")[1][:20])
            check("bold survives as strong", "<strong>" in h)
            check("figure carries its spec", 'data-spec=' in h)
            check("figure marked regenerable", 'data-regenerable="true"' in h)
            check("description box present", 'class="desc"' in h)
            check("machine index embedded", 'id="eduverse-book"' in h)
            check("entities not double escaped", "&amp;middot;" not in h)

            import re
            idx = json.loads(re.search(
                r'id="eduverse-book">(.*?)</script>', h, re.S).group(1))
            check("index has the outline", len(idx["outline"]) >= 3)
            check("index has the figures", len(idx["figures"]) >= 1)
            check("index has the equations",
                  any(e["latex"] for e in idx["equations"]))
            check("index has the tables", len(idx["tables"]) >= 1)
            check("index reports coverage",
                  idx["coverage"]["coverage"] == 1.0)
    finally:
        pl.OllamaClient = real


@group
def test_vlm_recovers_math():
    """An equation the text layer mangled must come back as LaTeX, not as a
    picture, and a formula the model is unsure of must never be typed out."""
    from eduverse_book2docx import figures as figmod

    class MathClient(_FakeClient):
        latex = "x^{2} + y^{2}"
        confident = True

        def read_math(self, png):
            return {"latex": self.latex, "confident": self.confident}

    cfg = Config()
    cfg.offline = False
    from PIL import Image
    crop = Image.new("RGB", (120, 40), "white")

    MathClient.confident = True
    got = figmod.read_math(MathClient(cfg), crop, cfg)
    check("confident formula accepted", got == "x^{2} + y^{2}", str(got))

    MathClient.confident = False
    check("unsure formula refused",
          figmod.read_math(MathClient(cfg), crop, cfg) is None)

    MathClient.confident = True
    MathClient.latex = "\\frac{a}{b"          # invalid
    check("invalid formula refused",
          figmod.read_math(MathClient(cfg), crop, cfg) is None)

    cfg.vlm_math = False
    MathClient.latex = "x"
    check("pass can be switched off",
          figmod.read_math(MathClient(cfg), crop, cfg) is None)


@group
def test_fragment_clustering():
    from eduverse_book2docx import figures as fm
    cfg = Config()
    area = 595 * 842

    # a clock: the face plus five small pieces sitting on top of it
    face = (300.0, 500.0, 420.0, 620.0)
    bits = [(340.0, 540.0, 372.0, 560.0), (352.0, 512.0, 366.0, 546.0),
            (356.0, 506.0, 368.0, 516.0), (400.0, 552.0, 412.0, 566.0),
            (306.0, 552.0, 320.0, 566.0)]
    groups = fm.cluster([face] + bits, [], area, cfg)
    check("clock comes out as one figure", len(groups) == 1, str(len(groups)))
    check("all six pieces accounted for",
          len(groups[0][1]) == 6, str(groups[0][1]))
    check("merged box covers the face",
          groups[0][0][0] <= 300.5 and groups[0][0][2] >= 419.5,
          str(groups[0][0]))

    # two figures side by side with a caption between them: must stay two
    a = (60.0, 500.0, 200.0, 600.0)
    b = (330.0, 500.0, 470.0, 600.0)
    text = [(210.0, 520.0, 320.0, 540.0)]
    groups = fm.cluster([a, b], text, area, cfg)
    check("text between keeps them apart", len(groups) == 2, str(len(groups)))

    # far apart, no text between: still two
    groups = fm.cluster([(60.0, 100.0, 160.0, 200.0),
                         (400.0, 600.0, 500.0, 700.0)], [], area, cfg)
    check("distant figures stay separate", len(groups) == 2)

    # a speck must not drag a big figure into a much bigger box
    big = (60.0, 100.0, 300.0, 340.0)
    speck = (520.0, 700.0, 528.0, 708.0)
    groups = fm.cluster([big, speck], [], area, cfg)
    check("a speck does not swallow the page", len(groups) == 2,
          str([g[0] for g in groups]))

    # touching pieces join
    groups = fm.cluster([(100.0, 100.0, 200.0, 200.0),
                         (200.0, 100.0, 300.0, 200.0)], [], area, cfg)
    check("touching pieces join", len(groups) == 1)

    # switched off, nothing merges
    off = Config()
    off.fragment_merge = False
    check("merging can be switched off",
          len(fm.cluster([face] + bits, [], area, off)) == 6)


@group
def test_image_hashing():
    from eduverse_book2docx import figures as fm
    from PIL import Image, ImageDraw
    cfg = Config()

    def circle(size, r, fill):
        im = Image.new("RGB", (size, size), "white")
        d = ImageDraw.Draw(im)
        d.ellipse([size // 2 - r, size // 2 - r, size // 2 + r, size // 2 + r],
                  fill=fill)
        return im

    a = circle(120, 40, "black")
    b = a.resize((90, 90)).resize((120, 120))     # same picture, requantised
    c = Image.new("RGB", (120, 120), "white")
    d = ImageDraw.Draw(c)
    d.rectangle([10, 60, 110, 110], fill="black")

    ha, hb, hc = (fm.phash(x, cfg.dup_hash_bits) for x in (a, b, c))
    check("same picture hashes near identical",
          fm.near_duplicate(ha, hb, cfg), str(fm.hamming(ha, hb)))
    check("different picture does not", not fm.near_duplicate(ha, hc, cfg),
          str(fm.hamming(ha, hc)))
    check("identical hashes have distance zero", fm.hamming(ha, ha) == 0)

    check("blank crop is flat", fm.is_flat(Image.new("RGB", (60, 60), "white")))
    check("real picture is not flat", not fm.is_flat(a))


@group
def test_duplicate_and_shattered_figures():
    """The two figure faults a real book produced: one clock arriving as six
    separate pictures, and a logo arriving once per page."""
    cfg = Config()
    cfg.offline = True
    with tempfile.TemporaryDirectory() as d:
        pdf = build_hard(os.path.join(d, "hard.pdf"))
        res = run(pdf, os.path.join(d, "out"), cfg)

        check("hard book balanced", res.ledger.balanced(),
              f"imbalance {res.ledger.imbalance()}")
        check("hard book gate passed", res.passed, "; ".join(res.failures))

        figs = [b for _, b in res.book.blocks() if b.type == "figure"]
        check("six part illustration is one figure",
              any("merged_from_6_fragments" in b.flags for b in figs),
              str([(b.id, b.flags) for b in figs]))
        check("no figure appears twice",
              len({(b.page, tuple(round(v) for v in b.bbox_pt)) for b in figs})
              == len(figs), str([b.id for b in figs]))

        # the logo printed on every page is gone
        check("repeated logo dropped",
              res.book.stats.get("repeating_images_dropped", 0) >= 6,
              str(res.book.stats.get("repeating_images_dropped")))
        check("watermark reason logged",
              any(e["reason"] == "watermark" for e in res.ledger.entries))
        check("no logo left in the book",
              all(b.width > 100 or b.height > 100 for b in figs),
              str([(b.id, round(b.width), round(b.height)) for b in figs]))

        # the same image placed twice at one spot is caught at source
        check("same placement deduped",
              any("placed twice" in (e.get("detail") or "")
                  for e in res.ledger.entries))

        # every fragment is still accounted for
        merged_drops = [e for e in res.ledger.entries
                        if "fragment merged" in (e.get("detail") or "")]
        check("merged fragments logged", len(merged_drops) == 5,
              str(len(merged_drops)))

        # figures on a page never overlap each other after clustering
        for page in res.book.pages:
            pf = [b for b in page.blocks if b.type == "figure"]
            for i, x in enumerate(pf):
                for y in pf[i + 1:]:
                    ix = min(x.x1, y.x1) - max(x.x0, y.x0)
                    iy = min(x.y1, y.y1) - max(x.y0, y.y0)
                    check(f"figures {x.id} and {y.id} do not overlap",
                          not (ix > 4 and iy > 4),
                          f"{x.bbox_pt} vs {y.bbox_pt}")


@group
def test_faults_from_the_real_output():
    """Four faults found by auditing a real ten page run: a table column
    emptied into the prose, a one column sidebar published as a table, a
    full page background published as a figure, and a shadow layer read as a
    stutter."""
    cfg = Config()
    cfg.offline = True
    with tempfile.TemporaryDirectory() as d:
        pdf = build_hard(os.path.join(d, "hard.pdf"))
        res = run(pdf, os.path.join(d, "out"), cfg)
        check("balanced", res.ledger.balanced(),
              f"imbalance {res.ledger.imbalance()}")
        check("gate passed", res.passed, "; ".join(res.failures))

        blocks = [b for _, b in res.book.blocks()]
        tbls = [b for b in blocks if b.type == "table"]
        si = next((t for t in tbls if t.table.cols == 3
                   and any(c.text == "Quantity" for c in t.table.cells)), None)
        check("SI table found", si is not None, str([t.table.cols for t in tbls]))
        if si:
            g = {(c.row, c.col): c.text for c in si.table.cells}
            check("first column is not empty",
                  all(g.get((r, 0), "").strip() for r in range(1, si.table.rows)),
                  str([g.get((r, 0)) for r in range(si.table.rows)]))
            check("labels land in the table, not the prose",
                  g.get((1, 0)) == "Amount of Substance", str(g.get((1, 0))))
            check("units still line up with their labels",
                  g.get((1, 1)) == "mole" and g.get((1, 2)) == "mol",
                  f"{g.get((1, 1))} {g.get((1, 2))}")

        prose = " ".join(b.text for b in blocks
                         if b.type in ("paragraph", "heading", "list_item"))
        check("table values did not leak into the prose",
              "Luminous Intensity" not in prose and "candela" not in prose,
              prose[-120:])

        check("no one column table",
              all(t.table.cols >= 2 for t in tbls),
              str([t.table.cols for t in tbls]))

        figs = [b for b in blocks if b.type == "figure"]
        page_area = res.book.pages[0].width_pt * res.book.pages[0].height_pt
        check("no page sized background figure",
              all(b.width * b.height < page_area * 0.9 for b in figs),
              str([(b.id, round(b.width), round(b.height)) for b in figs]))

        stutter = [b.text for b in blocks if "Analog" in (b.text or "")]
        check("shadow layer read once", stutter == ["Analog and digital clocks"],
              str(stutter))


@group
def test_text_dedupe():
    from eduverse_book2docx.structure import (_collapse_doubled, dedupe_text)
    cfg = Config()

    check("doubled words collapsed",
          _collapse_doubled("Analog clocks Analog clocks") == "Analog clocks")
    check("doubled phrase collapsed",
          _collapse_doubled("Unit Systems Unit Systems") == "Unit Systems")
    check("ordinary sentence untouched",
          _collapse_doubled("The angle of incidence equals the angle of "
                            "reflection.") is None)
    check("short text untouched", _collapse_doubled("ab ab") is None)
    check("repeated word is not doubled text",
          _collapse_doubled("very very good indeed") is None)

    page = Page(page=1, width_pt=595, height_pt=842)
    page.blocks = [
        Block(id="a", page=1, type="paragraph", bbox_pt=[56, 100, 300, 114],
              text="Measurement is the base of all study"),
        Block(id="b", page=1, type="paragraph", bbox_pt=[56.5, 100.4, 300, 114],
              text="Measurement is the base of all study"),
        Block(id="c", page=1, type="paragraph", bbox_pt=[56, 400, 300, 414],
              text="Measurement is the base of all study"),
    ]
    led = Ledger()
    led.emit(3)
    n = dedupe_text(page, led, cfg)
    check("shadow copy dropped", n == 1, str(n))
    check("same words elsewhere on the page kept",
          len(page.blocks) == 2, str([b.id for b in page.blocks]))
    check("emitted count corrected", led.emitted == 2)
    check("drop reason recorded",
          any("same text as" in (e.get("detail") or "") for e in led.entries))

    off = Config()
    off.dedupe_text = False
    page2 = Page(page=1, width_pt=595, height_pt=842)
    page2.blocks = [
        Block(id="a", page=1, type="paragraph", bbox_pt=[56, 100, 300, 114],
              text="the same words twice over here"),
        Block(id="b", page=1, type="paragraph", bbox_pt=[56, 100, 300, 114],
              text="the same words twice over here"),
    ]
    check("can be switched off",
          dedupe_text(page2, Ledger(), off) == 0)


@group
def test_table_cell_assignment():
    from eduverse_book2docx import tables as tmod
    spans = [
        {"text": "Amount of Substance", "bbox": (58.0, 440.0, 214.0, 452.0),
         "bold": False, "italic": False, "line_y": 440.0},
        {"text": "mole", "bbox": (211.0, 440.0, 240.0, 452.0),
         "bold": False, "italic": False, "line_y": 440.0},
    ]
    # the label overflows its cell on both sides, as printed labels do
    text, runs = tmod.cell_content(spans, (56.0, 436.0, 206.0, 458.0))
    check("overflowing label still lands in its cell",
          text == "Amount of Substance", repr(text))
    text2, _ = tmod.cell_content(spans, (206.0, 436.0, 316.0, 458.0))
    check("the next cell gets its own value", text2 == "mole", repr(text2))
    check("an empty cell stays empty",
          tmod.cell_content(spans, (400.0, 436.0, 500.0, 458.0))[0] == "")


@group
def test_equation_and_table_boxes():
    """Every equation and every table gets a visible description box in
    flow.docx, the same way a figure does."""
    cfg = Config()
    cfg.offline = True
    with tempfile.TemporaryDirectory() as d:
        pdf = build_sample(os.path.join(d, "s.pdf"))
        res = run(pdf, os.path.join(d, "out"), cfg)
        check("balanced", res.ledger.balanced(),
              f"imbalance {res.ledger.imbalance()}")
        check("gate passed", res.passed, "; ".join(res.failures))

        notes = {n.caption_of: n for _, n in res.book.blocks()
                 if n.type == "figure_note"}
        eqs = [b for _, b in res.book.blocks() if b.type == "equation"]
        tbls = [b for _, b in res.book.blocks() if b.type == "table"]

        # An equation gets NO description box by default: the formula is
        # already written out as a formula, and a box under it saying
        # "equals 1000, LaTeX: = 1000" is noise in the middle of the reading.
        check("no box for an equation by default",
              not any(e.id in notes for e in eqs), str(list(notes)))
        check("every table has a box",
              all(t.id in notes for t in tbls), str(list(notes)))

        # the box is still available behind its flag
        cfg_eq = Config()
        cfg_eq.offline = True
        cfg_eq.equation_description_box = True
        res_eq = run(pdf, os.path.join(d, "out_eq"), cfg_eq)
        notes = {n.caption_of: n for _, n in res_eq.book.blocks()
                 if n.type == "figure_note"}
        eqs = [b for _, b in res_eq.book.blocks() if b.type == "equation"]
        tbls = [b for _, b in res_eq.book.blocks() if b.type == "table"]
        check("the flag brings the equation box back",
              all(e.id in notes for e in eqs), str(list(notes)))

        en = notes[eqs[0].id]
        check("equation box is marked as one", en.block_kind == "equation")
        check("equation box reads the formula out",
              "divided by" in en.text, en.text[:80])
        check("equation box carries the latex verbatim",
              "LaTeX: " in en.text and eqs[0].latex in en.text, en.text[:120])
        check("equation box says where it came from",
              "Read from:" in en.text, en.text[:160])

        tn = notes[tbls[0].id]
        check("table box is marked as one", tn.block_kind == "table")
        check("table box gives the shape",
              "rows and" in tn.text and "columns" in tn.text, tn.text[:80])
        check("table box names the columns", "Columns:" in tn.text)
        check("table box writes the rows out",
              "225000" in tn.text, tn.text[:200])
        check("table box says whether it was ruled", "ruled" in tn.text)

        merged = next((t for t in tbls if any(c.col_span > 1
                                              for c in t.table.cells)), None)
        if merged:
            check("table box records the merge",
                  "Merged cells:" in notes[merged.id].text,
                  notes[merged.id].text[:200])

        # and they are visible in the docx
        from docx import Document
        doc = Document(res.files["flow_docx"])
        cells = " ".join(c.text for t in doc.tables
                         for r in t.rows for c in r.cells)
        check("no equation box in the docx by default",
              "Equation description" not in cells)
        check("equation box appears once the flag is on",
              "Equation description" in " ".join(
                  c.text for t in Document(res_eq.files["flow_docx"]).tables
                  for r in t.rows for c in r.cells))
        check("table box visible in the docx", "Table description" in cells)
        check("boxes sit after what they describe",
              _note_follows(res.book))

        # the digital book shows them too
        h = open(res.files["digital"], encoding="utf-8").read()
        check("no equation box in the digital book by default",
              "Equation description" not in h)
        check("table box in the digital book", "Table description" in h)

        # a round trip must not turn a box into a table
        r2 = run(res.files["flow_docx"], os.path.join(d, "out2"), cfg)
        c2 = r2.book.counts()
        check("boxes survive the round trip as boxes",
              c2.get("figure_note", 0) >= res.book.counts().get(
                  "figure_note", 0), str(c2))
        check("boxes did not become tables",
              c2.get("table", 0) == res.book.counts().get("table", 0),
              f"{res.book.counts().get('table')} -> {c2.get('table')}")


def _note_follows(book) -> bool:
    for page in book.pages:
        ids = [b.id for b in page.blocks]
        for b in page.blocks:
            if b.type != "figure_note" or b.caption_of not in ids:
                continue
            if ids.index(b.id) != ids.index(b.caption_of) + 1:
                return False
    return True


@group
def test_speak_formula():
    from eduverse_book2docx.mathtools import speak
    cases = {
        "n = \\frac{c}{v}": "divided by",
        "x^{2}": "squared",
        "y^{3}": "cubed",
        "\\sqrt{2}": "square root",
        "a \\times b": "times",
        "p \\geq q": "greater than or equal",
        "\\alpha + \\beta": "alpha",
        "v_{0}": "sub",
        "E = mc^{2}": "equals",
    }
    for tex, want in cases.items():
        got = speak(tex)
        check(f"speak {tex[:16]} says {want}", want in got, got)
    check("empty formula reads empty", speak("") == "")
    check("no latex backslashes leak through",
          "\\\\" not in speak("\\frac{\\alpha}{\\beta}"),
          speak("\\frac{\\alpha}{\\beta}"))


@group
def test_missing_description_is_visible():
    """A figure box that cannot be filled must say why. A silent gap reads
    like a missing feature; a stated failure reads like a failure."""
    from eduverse_book2docx import figures as figmod
    from eduverse_book2docx import pipeline as pl
    from eduverse_book2docx.ollama_client import OllamaError
    from PIL import Image

    class Dead(_FakeClient):
        def describe_figure(self, png):
            raise OllamaError("qwen3.5:4b: Connection refused")

        def figure_spec(self, png, kind):
            return {}

        def read_table(self, png):
            return {"is_table": False}

        def read_math(self, png):
            return {"latex": "", "confident": False}

    cfg = Config()
    cfg.offline = False
    meta = figmod.describe(Dead(cfg), Image.new("RGB", (80, 80), "white"), cfg)
    check("failure is recorded", "Connection refused" in meta.error, meta.error)
    check("no description invented", meta.desc == "")

    off = Config()
    off.describe_figures = False
    check("switched off is recorded",
          "switched off" in figmod.describe(
              Dead(off), Image.new("RGB", (80, 80), "white"), off).error)

    empty = Config()
    empty.offline = False

    class Silent(_FakeClient):
        def describe_figure(self, png):
            return {"fig_kind": "diagram", "desc": "   ", "entities": []}

        def figure_spec(self, png, kind):
            return {}

        def read_table(self, png):
            return {"is_table": False}

        def read_math(self, png):
            return {"latex": "", "confident": False}

    m2 = figmod.describe(Silent(empty), Image.new("RGB", (80, 80), "white"), empty)
    check("an empty answer is recorded too",
          "returned no description" in m2.error, m2.error)

    real = pl.OllamaClient
    pl.OllamaClient = Dead
    try:
        with tempfile.TemporaryDirectory() as d:
            pdf = build_sample(os.path.join(d, "s.pdf"))
            res = pl.run(pdf, os.path.join(d, "out"), cfg)

            check("balanced", res.ledger.balanced())
            notes = [n for _, n in res.book.blocks()
                     if n.type == "figure_note" and n.block_kind
                     not in ("equation", "table")]
            figs = [b for _, b in res.book.blocks() if b.type == "figure"]
            check("a box is still written for every figure",
                  len(notes) == len(figs), f"{len(notes)} of {len(figs)}")
            check("the box states the failure",
                  all("No description available" in n.text for n in notes),
                  notes[0].text[:80] if notes else "none")
            check("the box names the cause",
                  any("Connection refused" in n.text for n in notes))
            check("the box says what to run",
                  any("--check" in n.text for n in notes))

            check("no equation box, as configured",
                  not any(n.block_kind == "equation"
                          for _, n in res.book.blocks()
                          if n.type == "figure_note"))
            check("table box unaffected",
                  any(n.block_kind == "table" and "columns" in n.text
                      for _, n in res.book.blocks() if n.type == "figure_note"))

            check("coverage shows none described",
                  res.coverage.summary()["figures_described"] == 0)
            h = open(os.path.join(d, "out", "report.html"), encoding="utf-8").read()
            check("report warns about it",
                  "Figure descriptions are missing" in h)
            check("report names the cause", "Connection refused" in h)

            from docx import Document
            doc = Document(res.files["flow_docx"])
            cells = " ".join(c.text for t in doc.tables
                             for r in t.rows for c in r.cells)
            check("the docx shows the failure",
                  "No description available" in cells)
    finally:
        pl.OllamaClient = real


@group
def test_describe_one_command():
    """--describe proves the model works on one image, so a missing
    description can be pinned on the model or on the pipeline, not guessed."""
    from PIL import Image, ImageDraw
    from eduverse_book2docx import cli

    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "clock.png")
        im = Image.new("RGB", (240, 240), "white")
        dr = ImageDraw.Draw(im)
        dr.ellipse([20, 20, 220, 220], outline="black", width=3)
        im.save(path)

        class Good(OllamaClient):
            def health(self):
                return {"ok": True, "models": [self.cfg.vlm_model]}

            def describe_figure(self, png):
                return {"fig_kind": "diagram",
                        "desc": "A circle drawn in outline.",
                        "entities": ["circle"], "relationships": [],
                        "data_points": [], "text_in_image": []}

        class Dead(Good):
            def health(self):
                return {"ok": False, "error": "Connection refused"}

        class Silent(Good):
            def describe_figure(self, png):
                return {"fig_kind": "diagram", "desc": ""}

        real = cli.OllamaClient
        try:
            cli.OllamaClient = Good
            check("a working model returns success",
                  cli._describe_one(path, Config()) == 0)
            cli.OllamaClient = Dead
            check("an unreachable model returns failure",
                  cli._describe_one(path, Config()) == 1)
            cli.OllamaClient = Silent
            check("an empty answer returns failure",
                  cli._describe_one(path, Config()) == 1)
            cli.OllamaClient = Good
            check("a missing file returns 2",
                  cli._describe_one(os.path.join(d, "nope.png"), Config()) == 2)
        finally:
            cli.OllamaClient = real

    a = cli.build_parser().parse_args(["--describe", "x.png"])
    check("flag is parsed", a.describe == "x.png")
    check("no input file needed with the flag", a.input is None)


@group
def test_layout_engine():
    """Rules A and B: layout before reading."""
    from eduverse_book2docx import layout as L
    cfg = Config()

    # two columns of lines
    two = []
    for i in range(12):
        two.append((56.0, 100.0 + i * 14, 250.0, 112.0 + i * 14))
        two.append((330.0, 100.0 + i * 14, 520.0, 112.0 + i * 14))
    cols = L.columns_in(two, 56, 540, 90, 300, cfg)
    check("two columns found", len(cols) == 2, str(cols))

    # a stray mark in the gutter must not hide it
    noisy = two + [(325.0, 200.0, 330.0, 206.0)]
    check("stray mark in the gutter ignored",
          len(L.columns_in(noisy, 56, 540, 90, 300, cfg)) == 2)

    one = [(56.0, 100.0 + i * 14, 520.0, 112.0 + i * 14) for i in range(12)]
    check("one column stays one",
          len(L.columns_in(one, 56, 540, 90, 300, cfg)) == 1)

    # a ragged right margin is not a column
    ragged = [(56.0, 100.0 + i * 14, 400.0 + (i % 3) * 30, 112.0 + i * 14)
              for i in range(12)]
    check("ragged margin is not a column",
          len(L.columns_in(ragged, 56, 540, 90, 300, cfg)) == 1)

    # too few lines: do not guess
    check("too few lines means one column",
          len(L.columns_in(two[:4], 56, 540, 90, 300, cfg)) == 1)

    # bands only exist when there are columns
    lay = L.analyse(one, (0, 0, 595, 842), [56, 56, 56, 56], [], cfg)
    check("single column page is one band", len(lay.bands) == 1, str(lay.bands))
    check("wrapped lines do not create bands", lay.max_columns() == 1)

    # a full width heading over two columns splits the page into bands
    banded = [(56.0, 60.0, 520.0, 80.0)] + two
    lay2 = L.analyse(banded, (0, 0, 595, 842), [56, 56, 56, 56], [], cfg)
    check("full width line makes a band", len(lay2.bands) >= 2,
          str(len(lay2.bands)))
    check("columns still found under it", lay2.max_columns() == 2)


@group
def test_boxes_and_badges():
    from eduverse_book2docx import layout as L
    cfg = Config()
    panel = {"rect": (60.0, 100.0, 520.0, 320.0), "fill": (1, 1, 1),
             "color": None, "width": None}
    badge = {"rect": (70.0, 180.0, 150.0, 250.0), "fill": None,
             "color": (0, 0, 0), "width": 1.0}
    lines = [(180.0, 110.0, 500.0, 124.0),      # text above the badge
             (90.0, 205.0, 140.0, 222.0),       # the badge word
             (180.0, 205.0, 500.0, 222.0),      # text wrapping beside it
             (90.0, 260.0, 500.0, 276.0)]
    boxes = L.find_boxes([panel, badge], (0, 0, 595, 842), lines, cfg)
    check("panel and badge both found", len(boxes) == 2, str(len(boxes)))
    parent = next((b for b in boxes if b.kind != "badge"), None)
    check("badge marked as such",
          any(b.kind == "badge" for b in boxes), str([b.kind for b in boxes]))
    check("panel knows its title line",
          parent is not None and parent.title_line == 1,
          str(parent.title_line if parent else None))

    # a box holding no text is not a box
    empty = L.find_boxes([panel], (0, 0, 595, 842), [], cfg)
    check("an empty region is not a box", empty == [])


@group
def test_line_joining_rules():
    from eduverse_book2docx import layout as L, lines as LN
    cfg = Config()

    def mk(text, x0, y0, size=10.5, bold=0.0):
        return LN.Line(bbox=(x0, y0, x0 + 200, y0 + 12), text=text,
                       max_size=size, bold_frac=bold)

    def run(ls):
        lay = L.analyse([l.bbox for l in ls], (0, 0, 595, 842),
                        [56, 56, 56, 56], [], cfg)
        return LN.group_lines(ls, lay, 595.0, cfg)

    g = run([mk("the speed of light in", 56, 100),
             mk("a vacuum is constant", 56, 113)])
    check("wrapped sentence joins", len(g) == 1, str(len(g)))

    g = run([mk("broken hyphen-", 56, 100), mk("ated word", 56, 113)])
    check("hyphenated word joins", len(g) == 1)

    g = run([mk("Length", 56, 100), mk("Mass", 56, 113),
             mk("Time", 56, 126), mk("Temperature", 56, 139)])
    check("a list of terms stays apart", len(g) == 4, str(len(g)))

    g = run([mk("Learning Objectives", 56, 100, bold=1.0),
             mk("After this lesson students will", 56, 113)])
    check("a bold label stays its own block", len(g) == 2, str(len(g)))

    g = run([mk("small print here", 56, 100, size=8.0),
             mk("body text here", 56, 113, size=10.5)])
    check("a size change breaks the join", len(g) == 2)


@group
def test_subject_detection():
    from eduverse_book2docx.sources.pdf_source import _find_subject
    check("a state name is never the subject",
          _find_subject("government of tamil nadu textbook", None) is None)
    check("the real subject is found beside the state",
          _find_subject("government of tamil nadu standard eight science",
                        None) == "Science")
    check("the most specific subject wins",
          _find_subject("samacheer kalvi social science class 8", None)
          == "Social Science")
    check("no subject means none",
          _find_subject("a book about nothing in particular", None) is None)


@group
def test_table_cell_recovery():
    """Rules 35 to 38."""
    from eduverse_book2docx import tables as tmod
    cfg = Config()
    td_cells = []
    from eduverse_book2docx.schema import TableCell, TableData
    for r in range(3):
        for c in range(2):
            td_cells.append(TableCell(row=r, col=c,
                                      text="" if (r, c) == (1, 0) else f"v{r}{c}"))
    td = TableData(rows=3, cols=2, cells=td_cells, header_rows=1)
    gaps = tmod.row_gaps(td)
    check("part empty row is flagged", gaps == [1], str(gaps))

    full = TableData(rows=2, cols=2, header_rows=1, cells=[
        TableCell(row=0, col=0, text="a"), TableCell(row=0, col=1, text="b"),
        TableCell(row=1, col=0, text="c"), TableCell(row=1, col=1, text="d")])
    check("a complete table has no gaps", tmod.row_gaps(full) == [])


@group
def test_figure_merge_guards():
    """Rules 26 to 29."""
    from eduverse_book2docx import figures as fm
    cfg = Config()
    area = 595 * 842
    a = (60.0, 500.0, 200.0, 600.0)
    b = (204.0, 500.0, 350.0, 600.0)

    check("touching pieces merge",
          len(fm.cluster([a, b], [], area, cfg)) == 1)
    check("a caption between them keeps them apart",
          len(fm.cluster([a, b], [], area, cfg,
                         captions=[(199.0, 540.0, 206.0, 556.0)])) == 2)
    check("a column gutter keeps them apart",
          len(fm.cluster([a, b], [], area, cfg,
                         gutters=[(200.0, 204.0)])) == 2)
    check("a panel edge keeps them apart",
          len(fm.cluster([a, b], [], area, cfg,
                         box_edges=[(50.0, 490.0, 202.0, 610.0)])) == 2)

    tiles = [(60.0, 500.0, 160.0, 600.0), (165.0, 500.0, 265.0, 600.0),
             (60.0, 605.0, 160.0, 705.0), (165.0, 605.0, 265.0, 705.0)]
    check("a tidy grid is one figure", fm.is_grid(tiles, cfg))
    check("two odd shapes are not a grid",
          not fm.is_grid([(0.0, 0.0, 10.0, 10.0), (0.0, 0.0, 300.0, 20.0)], cfg))


@group
def test_no_generation():
    """Nothing may be published that is not on the page."""
    cfg = Config()
    cfg.offline = True
    with tempfile.TemporaryDirectory() as d:
        pdf = build_sample(os.path.join(d, "s.pdf"))
        res = run(pdf, os.path.join(d, "out"), cfg)
        source = " ".join(
            b.text for _, b in res.book.blocks() if b.text).lower()
        for invented in ("in summary", "as we have seen", "note that the",
                         "this chapter", "let us now"):
            check(f"nothing generated: {invented!r}", invented not in source)
        # every figure description is empty offline, never invented
        for _, b in res.book.blocks():
            if b.type == "figure" and b.figure:
                check("no invented description offline",
                      b.figure.desc == "", b.figure.desc[:40])


# ---------------------------------------------------------------------------
# B14: Phase 1 review fixes
# ---------------------------------------------------------------------------
@group
def test_b14_colours_and_fonts():
    from eduverse_book2docx.textclean import font_family, readable_color
    check("white text is not written white", readable_color("FFFFFF") is None)
    check("pale text is not written", readable_color("F4F2EC") is None)
    check("accent colour kept", readable_color("905500") == "905500")
    check("family drops weight", font_family("ABC+Verdana-Bold") == "verdana")
    check("family joins spaced name",
          font_family("Minion Pro Bold") == font_family("MinionPro-Regular"))


@group
def test_b14_display_labels_are_headings():
    from eduverse_book2docx.config import Config
    from eduverse_book2docx.schema import Block, Page, Run
    from eduverse_book2docx.structure import SizeProfile, classify_text_block
    cfg = Config()
    prof = SizeProfile(cfg)
    prof.observe(12.0, 5000, 1)
    prof.observe_font("minion", 5000)
    prof.finalise()
    page = Page(page=5, width_pt=654, height_pt=881,
                margins_pt=[86, 80, 86, 80])

    def blk(text, font, color=None, bold=True, size=12.0):
        b = Block(id="x", page=5, type="paragraph", bbox_pt=[95, 400, 220, 417],
                  text=text, runs=[Run(s=0, e=len(text), bold=bold,
                                       color=color, font=font)])
        classify_text_block(b, page, prof, cfg, size, bold, color)
        return b
    b = blk("Learning Objectives", "Verdana-Bold", "FFFFFF")
    check("knockout label is a heading", b.type == "heading", b.type)
    check("knockout label sits at level 2", b.level == 2, str(b.level))
    b = blk("Activity 1", "Verdana-Bold", "905500")
    check("display face label is a heading", b.type == "heading", b.type)
    b = blk("1. Analog clocks", "Verdana-Bold")
    check("numbered label in display face is a heading", b.type == "heading",
          b.type)
    b = blk("Measure the length of your book.", "Verdana-Bold")
    check("a sentence is not a label", b.type != "heading", b.type)
    b = blk("Important words", "MinionPro-Bold")
    check("bold body face alone is not promoted", b.type != "heading", b.type)


@group
def test_b14_numbered_captions_only():
    from eduverse_book2docx.structure import is_numbered_caption
    pf = ("fig", "figure", "table")
    check("Table of Contents is not a caption",
          not is_numbered_caption("table of contents", pf))
    check("Table 1.2 is a caption", is_numbered_caption("table 1.2 units", pf))
    check("Fig. 1.13 is a caption", is_numbered_caption("fig. 1.13 x", pf))


@group
def test_b14_late_block_placement():
    from eduverse_book2docx.schema import Block, Page
    from eduverse_book2docx.placement import place_late_blocks
    pg = Page(page=2, width_pt=654, height_pt=881, columns=1,
              margins_pt=[86, 80, 86, 80])
    t = lambda i, y, x=98: Block(id=i, page=2, type="paragraph",
                                 bbox_pt=[x, y, x + 100, y + 15], text=i,
                                 source="native_span")
    stray = t("label", 523, 121)          # printed on the seal
    seal = Block(id="seal", page=2, type="figure", bbox_pt=[98, 487, 154, 544])
    pg.blocks = [stray, t("gov", 98), t("sale", 379), t("content", 456),
                 t("scert", 547), seal]
    place_late_blocks(pg)
    ids = [b.id for b in pg.blocks]
    check("a stray block does not drag the seal to the top",
          ids.index("seal") == ids.index("content") + 1, str(ids))

    pg2 = Page(page=3, width_pt=654, height_pt=881, columns=2,
               margins_pt=[86, 80, 86, 80])
    title = t("title", 77, 210)
    table = Block(id="table", page=3, type="table", bbox_pt=[88, 132, 568, 563])
    pg2.blocks = [title, table]
    place_late_blocks(pg2)
    check("a full width table stays under the page title",
          [b.id for b in pg2.blocks] == ["title", "table"],
          str([b.id for b in pg2.blocks]))

    pg3 = Page(page=3, width_pt=654, height_pt=881, columns=2,
               margins_pt=[86, 80, 86, 80])
    qrs = [Block(id=f"qr{i}", page=3, type="figure",
                 bbox_pt=[x, 576, x + 78, 663])
           for i, x in enumerate((134, 286, 442))]
    pg3.blocks = [t("note", 707, 126)] + list(reversed(qrs))
    place_late_blocks(pg3)
    check("a row of codes reads left to right before the text below",
          [b.id for b in pg3.blocks] == ["qr0", "qr1", "qr2", "note"],
          str([b.id for b in pg3.blocks]))


@group
def test_b14_side_heading_and_labels():
    from eduverse_book2docx.config import Config
    from eduverse_book2docx.schema import Block, Page
    from eduverse_book2docx.placement import hoist_side_headings
    from eduverse_book2docx.structure import bind_figure_labels
    pg = Page(page=4, width_pt=654, height_pt=881, columns=1)
    body = Block(id="b1", page=4, type="paragraph", bbox_pt=[134, 345, 346, 381],
                 text="This Science Textbook has 23 units.")
    body2 = Block(id="b2", page=4, type="paragraph",
                  bbox_pt=[134, 384, 451, 467], text="Each unit has activities.")
    head = Block(id="h", page=4, type="heading", level=3,
                 bbox_pt=[481, 360, 566, 418], text="HOW TO USE THE BOOK")
    pg.blocks = [body, body2, head]
    hoist_side_headings(pg)
    check("side heading comes before the text beside it",
          pg.blocks[0].id == "h", str([b.id for b in pg.blocks]))

    pg = Page(page=3, width_pt=654, height_pt=881, columns=1)
    qr = Block(id="qr", page=3, type="figure", bbox_pt=[134, 576, 212, 663])
    lab = Block(id="lab", page=3, type="paragraph", bbox_pt=[144, 671, 202, 693],
                text="E - book")
    pg.blocks = [qr, lab]
    bind_figure_labels(pg, Config())
    check("a word under a picture is its caption",
          lab.type == "caption" and lab.caption_of == "qr")


@group
def test_b14_glossary_rows():
    from eduverse_book2docx.lines import Group, Line, pair_label_rows
    terms, defs = [], []
    for i, (term, d) in enumerate([("Calibration", "Configuring an instrument."),
                                   ("Electric current", "Flow of charges."),
                                   ("Plane angle", "Angle of two lines."),
                                   ("Solid angle", "Angle of three planes.")]):
        y = 132 + 20 * i
        terms.append(Group(lines=[Line(bbox=(91, y, 150, y + 15), text=term,
                                       bold_frac=1.0)]))
        defs.append(Line(bbox=(201, y, 450, y + 15), text=d, bold_frac=0.0))
    groups = terms + [Group(lines=defs, column=1)]
    out = pair_label_rows(groups)
    texts = [" ".join(l.text for l in g.lines) for g in out]
    check("each term keeps its own definition", len(out) == 4, str(texts))
    check("row reads term then definition",
          texts[1] == "Electric current Flow of charges.", str(texts))


@group
def test_b14_positional_furniture():
    from eduverse_book2docx.config import Config
    from eduverse_book2docx.ledger import Ledger
    from eduverse_book2docx.schema import Block, Book, Page
    from eduverse_book2docx.structure import drop_furniture
    book = Book(book_id="x", source_file="x.pdf")
    for n in range(1, 21):
        p = Page(page=n, width_pt=654, height_pt=881)
        p.blocks.append(Block(id=f"p{n}b1", page=n, type="paragraph",
                              bbox_pt=[86, 300, 316, 400], text="Body text."))
        if n in (7, 9, 11, 13):
            p.blocks.append(Block(id=f"p{n}f", page=n, type="paragraph",
                                  bbox_pt=[519, 809, 565, 820],
                                  text="Measurement"))
        book.pages.append(p)
    led = Ledger()
    for _ in range(24):
        led.detect(); led.emit()
    drop_furniture(book, led, Config())
    left = [b.text for p in book.pages for b in p.blocks]
    check("a unit footer on four pages is furniture",
          "Measurement" not in left, str(left.count("Measurement")))
    check("body text kept", left.count("Body text.") == 20)


@group
def test_b14_captioned_figure_is_not_a_table():
    class _TableFake(_FakeClient):
        def read_table(self, png):
            return {"is_table": True, "header": ["Ammeter", "Battery", "Lamp"],
                    "rows": [["A", "", ""], ["x", "y", "z"]]}
    from eduverse_book2docx import pipeline
    desc = {"fig_kind": "diagram", "text_only": False, "label": "",
            "desc": "A circuit.", "entities": [], "relationships": [],
            "data_points": [], "text_in_image": []}
    with tempfile.TemporaryDirectory() as d:
        cfg = Config()
        cfg.offline = False
        _FakeClient.describe_as = desc
        _FakeClient.ocr_markdown = ""
        real = pipeline.OllamaClient
        pipeline.OllamaClient = _TableFake
        try:
            pdf = build_sample(os.path.join(d, "s.pdf"))
            res = pipeline.run(pdf, os.path.join(d, "out"), cfg)
        finally:
            pipeline.OllamaClient = real
        vlm_tables = [b for _, b in res.book.blocks()
                      if b.type == "table" and b.source == "vlm_table"]
        check("a diagram is never offered as a table", not vlm_tables,
              str(len(vlm_tables)))
        check("the diagram stays a figure",
              res.book.counts().get("figure", 0) >= 1)


# ---------------------------------------------------------------------------
def main() -> int:
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for t in tests:
        t()
    print(f"\n{len(PASS)} passed, {len(FAIL)} failed\n")
    for name, detail in FAIL:
        print(f"FAIL  {name}")
        if detail:
            print("      " + detail.strip().replace("\n", "\n      ")[:600])
    return 1 if FAIL else 0


if __name__ == "__main__":
    raise SystemExit(main())
