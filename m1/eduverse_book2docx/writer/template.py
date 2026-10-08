"""Builds eduverse_base.docx, the style sheet every output inherits.

Word is fussy about styles created on the fly, so the template is built once
and then loaded with Document(template_path) for every run.
"""

from __future__ import annotations

import os

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..config import Config


def build(path: str, cfg: Config) -> str:
    doc = Document()

    base = doc.styles["Normal"]
    base.font.name = cfg.styles["EV Body"][0]
    base.font.size = Pt(cfg.styles["EV Body"][1])
    _east_asia(base, cfg.styles["EV Body"][0])

    for name, (font, size, bold, before, after) in cfg.styles.items():
        try:
            st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        except ValueError:
            st = doc.styles[name]
        st.font.name = font
        st.font.size = Pt(size)
        st.font.bold = bold
        _east_asia(st, font)
        pf = st.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = name.startswith("EV H") or name == "EV Title"
        if name.startswith("EV H") or name == "EV Title":
            st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            _outline_level(st, _level_of(name))
        if name == "EV Caption":
            st.font.italic = True
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if name == "EV Equation":
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if name == "EV Note":
            st.font.color.rgb = RGBColor(0x8A, 0x8A, 0x8A)
        if name.startswith("EV List"):
            pf.left_indent = Pt(18)
            # numbering lives in the built-in list styles; never insert a
            # bullet glyph as literal text
            builtin = "List Number" if "Number" in name else "List Bullet"
            _copy_numbering(doc, builtin, st)

    # a character style for hidden machine text
    try:
        hid = doc.styles.add_style("EV Hidden", WD_STYLE_TYPE.CHARACTER)
        hid.font.hidden = True
        hid.font.size = Pt(1)
    except ValueError:
        pass

    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    _fix_settings(doc)
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    doc.save(path)
    return path


def _fix_settings(doc) -> None:
    """The stock python-docx settings.xml ships a w:zoom with no percent,
    which fails XSD validation."""
    try:
        settings = doc.settings.element
    except AttributeError:
        return
    for z in settings.findall(qn("w:zoom")):
        if not z.get(qn("w:percent")):
            z.set(qn("w:percent"), "100")


def _copy_numbering(doc, builtin_name: str, target) -> None:
    try:
        src = doc.styles[builtin_name]
    except KeyError:
        return
    src_ppr = src.element.find(qn("w:pPr"))
    if src_ppr is None:
        return
    numpr = src_ppr.find(qn("w:numPr"))
    if numpr is None:
        return
    import copy
    from .ooxml import PPR_ORDER, replace_ordered
    replace_ordered(target.element.get_or_add_pPr(), copy.deepcopy(numpr),
                    PPR_ORDER)


def _level_of(name: str) -> int:
    if name == "EV Title":
        return 0
    try:
        return int(name.replace("EV H", ""))
    except ValueError:
        return 9


def _outline_level(style, level: int) -> None:
    """Without an outlineLvl a custom heading style never reaches the TOC."""
    from .ooxml import PPR_ORDER, replace_ordered
    el = style.element.makeelement(qn("w:outlineLvl"), {})
    el.set(qn("w:val"), str(max(0, min(8, level))))
    replace_ordered(style.element.get_or_add_pPr(), el, PPR_ORDER)


def _east_asia(style, font: str) -> None:
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font)


def ensure(out_dir: str, cfg: Config) -> str:
    path = os.path.join(out_dir, cfg.template_name)
    if not os.path.exists(path):
        build(path, cfg)
    return path
